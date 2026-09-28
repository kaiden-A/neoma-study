"""File text extraction and search inside files."""

import io
import uuid
from typing import Any, cast

from fastapi.testclient import TestClient
from sqlalchemy.orm import Session as DbSession

from app.models import FileObject

PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _set_text(shape: object, text: str) -> None:
    cast(Any, shape).text = text


def _pptx_bytes() -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    layout = presentation.slide_layouts[1]
    slide = presentation.slides.add_slide(layout)
    _set_text(slide.shapes.title, "Photosynthesis")
    _set_text(slide.placeholders[1], "chloroplasts capture light")
    second = presentation.slides.add_slide(layout)
    _set_text(second.shapes.title, "Respiration")
    _set_text(second.placeholders[1], "mitochondria release energy")
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_upload_extracts_and_search_reports_the_page(
    client: TestClient, sign_in, make_user, db: DbSession, storage
) -> None:
    sign_in(make_user(db))
    upload = client.post(
        "/api/files", files={"file": ("lecture.pptx", io.BytesIO(_pptx_bytes()), PPTX_TYPE)}
    ).json()

    row = db.get(FileObject, uuid.UUID(upload["id"]))
    assert row is not None and row.extracted_text
    assert "chloroplasts capture light" in row.extracted_text
    assert row.extracted_at is not None
    note = client.post(
        "/api/notes", json={"type": "slides", "title": "Week 4", "fileId": upload["id"]}
    ).json()

    results = client.get("/api/notes?q=chloroplasts").json()

    assert [item["id"] for item in results] == [note["id"]]
    assert results[0]["matchSnippet"].startswith("found in file · page 1:")
    assert "chloroplasts" in results[0]["matchSnippet"]

    second_slide = client.get("/api/notes?q=mitochondria").json()
    assert "page 2" in second_slide[0]["matchSnippet"]


def test_search_still_matches_text_fields_and_escapes_wildcards(
    client: TestClient, sign_in, make_user, db: DbSession
) -> None:
    sign_in(make_user(db))
    tag_note = client.post(
        "/api/notes", json={"title": "Redox", "body": "titration", "tags": ["chem-301"]}
    ).json()
    client.post("/api/notes", json={"title": "Recursion", "body": "call stacks"})

    assert [item["id"] for item in client.get("/api/notes?q=chem-301").json()] == [tag_note["id"]]
    assert [item["id"] for item in client.get("/api/notes?q=redox").json()] == [tag_note["id"]]
    # A wildcard is a literal, not a pattern.
    assert client.get("/api/notes?q=%25").json() == []
    assert client.get("/api/notes?q=_").json() == []


def test_group_notes_are_searchable_by_members(
    client: TestClient, sign_in, make_user, db: DbSession
) -> None:
    owner = make_user(db, email="ada@example.com")
    sign_in(owner)
    group = client.post(
        "/api/groups",
        json={"kind": "study", "name": "Finals crew", "subject": "", "color": "violet", "description": ""},
    ).json()
    note = client.post(
        f"/api/groups/{group['id']}/notes", json={"title": "Group gem", "body": "secret sauce"}
    ).json()

    results = client.get("/api/notes?q=secret").json()

    assert [item["id"] for item in results] == [note["id"]]


def test_binary_files_are_left_alone(
    client: TestClient, sign_in, make_user, db: DbSession, storage
) -> None:
    sign_in(make_user(db))
    upload = client.post(
        "/api/files", files={"file": ("scan.jpg", io.BytesIO(b"not-an-image"), "image/jpeg")}
    ).json()

    row = db.get(FileObject, uuid.UUID(upload["id"]))
    assert row is not None
    assert row.extracted_text is None
    assert client.get("/api/notes?q=not-an-image").json() == []
