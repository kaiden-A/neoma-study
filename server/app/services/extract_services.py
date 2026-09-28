"""Best-effort text extraction for uploaded files.

Searching inside decks and past papers is what makes the vault useful; the
upload itself must never fail because a parser choked, so the entry point
swallows its errors and logs instead. PDF and slide text carries
`--- page N ---` markers so a search hit can name the page.
"""

import io
import logging
import re

from sqlalchemy.orm import Session as DbSession

from ..models import FileObject, utcnow
from . import storage_services

log = logging.getLogger(__name__)

MAX_EXTRACT_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 500_000
PAGE_MARKER = "--- page {page} ---"
PAGE_PATTERN = re.compile(r"--- page (\d+) ---")

DOCX_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
}
PPTX_TYPES = {
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "application/vnd.ms-powerpoint",
}


def supports(content_type: str) -> bool:
    return (
        content_type == "application/pdf"
        or content_type in DOCX_TYPES
        or content_type in PPTX_TYPES
        or content_type.startswith("text/")
        or content_type == "application/json"
    )


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        pages.append(f"{PAGE_MARKER.format(page=index)}\n{page.extract_text() or ''}")
    return "\n\n".join(pages).strip()


def _docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    return "\n".join(paragraph.text for paragraph in document.paragraphs).strip()


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    chunks = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines = []
        for shape in slide.shapes:
            text = getattr(shape, "text", None)
            if isinstance(text, str) and text:
                lines.append(text)
        chunks.append(f"{PAGE_MARKER.format(page=index)}\n" + "\n".join(lines))
    return "\n\n".join(chunks).strip()


def extract(content_type: str, data: bytes) -> str | None:
    if content_type == "application/pdf":
        return _pdf(data)
    if content_type in DOCX_TYPES:
        return _docx(data)
    if content_type in PPTX_TYPES:
        return _pptx(data)
    if content_type.startswith("text/") or content_type == "application/json":
        return data.decode("utf-8", errors="replace").strip()
    return None


def search_page(text: str, index: int) -> int | None:
    """The last page marker before a character index, or None."""
    page = None
    for match in PAGE_PATTERN.finditer(text[:index]):
        page = int(match.group(1))
    return page


def extract_file_text(db: DbSession, storage: storage_services.Storage, record: FileObject) -> bool:
    """Reads the object, extracts and stores its text. Never raises."""
    if not supports(record.content_type) or record.size > MAX_EXTRACT_BYTES:
        return False
    try:
        data = storage.get_bytes(record.key, MAX_EXTRACT_BYTES)
        if data is None:
            return False
        text = extract(record.content_type, data)
        if text is None:
            return False
        record.extracted_text = text[:MAX_TEXT_CHARS] or None
        record.extracted_at = utcnow()
        db.commit()
        return True
    except Exception:  # noqa: BLE001 - a parser must never fail the upload
        log.warning("text extraction failed for %s", record.key, exc_info=True)
        db.rollback()
        return False
